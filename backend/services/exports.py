"""CSV / Excel / PDF report exports (Q15 Better Reporting).

CSV and Excel cells are protected against formula injection: values starting
with = + - @ (or tab/CR) are prefixed with an apostrophe.
"""
from __future__ import annotations

import csv
import io

from flask import Response
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..utils import fmt_dt, utcnow

DANGEROUS = ("=", "+", "-", "@", "\t", "\r")


def _safe(value):
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    if text.startswith(DANGEROUS):
        # keep plain negative numbers usable
        try:
            float(text)
            return text
        except ValueError:
            return "'" + text
    return text


def _filename(base: str, ext: str) -> str:
    return f"{base}-{utcnow().strftime('%Y%m%d-%H%M')}.{ext}"


def csv_response(base: str, headers: list[str], rows: list[list]) -> Response:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for r in rows:
        writer.writerow([_safe(v) for v in r])
    return Response("\ufeff" + buf.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": f"attachment; filename={_filename(base, 'csv')}"})


def xlsx_response(base: str, title: str, headers: list[str], rows: list[list]) -> Response:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append([f"QuizMaster — {title}"])
    ws["A1"].font = Font(bold=True, size=14, color="312E81")
    ws.append([f"Generated {fmt_dt(utcnow())}"])
    ws["A2"].font = Font(italic=True, color="64748B")
    ws.append([])
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="6366F1")
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=4, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([_safe(v) for v in r])
    for col in range(1, len(headers) + 1):
        width = max([len(str(headers[col - 1]))] + [len(str(r[col - 1])) for r in rows[:200] if col - 1 < len(r)])
        ws.column_dimensions[get_column_letter(col)].width = min(48, max(10, width + 2))
    ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename={_filename(base, 'xlsx')}"})


def _latin(text) -> str:
    return str("" if text is None else text).encode("latin-1", "replace").decode("latin-1")


def pdf_response(base: str, title: str, headers: list[str], rows: list[list], widths: list[float] | None = None,
                 summary: list[tuple[str, str]] | None = None) -> Response:
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_fill_color(99, 102, 241)
    pdf.rect(0, 0, 297, 22, "F")
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 16)
    pdf.set_xy(10, 6)
    pdf.cell(0, 10, _latin(f"QuizMaster  |  {title}"))
    pdf.set_text_color(30, 27, 75)
    pdf.set_font("Helvetica", "", 9)
    pdf.set_xy(10, 26)
    pdf.cell(0, 6, _latin(f"Generated {fmt_dt(utcnow())}  -  {len(rows)} record(s)"))
    pdf.ln(8)
    if summary:
        pdf.set_font("Helvetica", "B", 9)
        for k, v in summary:
            pdf.cell(60, 6, _latin(k))
            pdf.set_font("Helvetica", "", 9)
            pdf.cell(0, 6, _latin(v), new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", "B", 9)
        pdf.ln(3)
    usable = 277
    widths = widths or [usable / len(headers)] * len(headers)
    scale = usable / sum(widths)
    widths = [w * scale for w in widths]

    def header():
        pdf.set_font("Helvetica", "B", 8.5)
        pdf.set_fill_color(238, 242, 255)
        for h, w in zip(headers, widths):
            pdf.cell(w, 7, _latin(h), border=0, fill=True)
        pdf.ln()

    header()
    pdf.set_font("Helvetica", "", 8)
    for i, r in enumerate(rows):
        if pdf.get_y() > 190:
            pdf.add_page()
            header()
            pdf.set_font("Helvetica", "", 8)
        if i % 2:
            pdf.set_fill_color(248, 250, 252)
        else:
            pdf.set_fill_color(255, 255, 255)
        for v, w in zip(r, widths):
            text = _latin(v)
            max_chars = max(4, int(w / 1.75))
            if len(text) > max_chars:
                text = text[: max_chars - 1] + "~"
            pdf.cell(w, 6.2, text, border=0, fill=True)
        pdf.ln()
    out = bytes(pdf.output())
    return Response(out, mimetype="application/pdf",
                    headers={"Content-Disposition": f"attachment; filename={_filename(base, 'pdf')}"})


def export(fmt: str, base: str, title: str, headers: list[str], rows: list[list], **kw) -> Response:
    if fmt == "xlsx":
        return xlsx_response(base, title, headers, rows)
    if fmt == "pdf":
        return pdf_response(base, title, headers, rows, **kw)
    return csv_response(base, headers, rows)
