"""Design system + low-level helpers for the QuizMaster deck.

Handles the things python-pptx cannot do natively, via OOXML:
  * gradient fills (shapes and text runs) in the QuizMaster brand colours
  * soft drop shadows, rounded-corner pictures
  * slide transitions (manual advance - click to move on, never auto-play)
  * gentle staggered fade-in entrance animations (play on slide entry)
  * hanging-indent bullet lists, browser / phone screenshot frames
"""
from lxml import etree
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- palette ---
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
BODY    = RGBColor(0xCD, 0xD1, 0xF2)
MUTED   = RGBColor(0x96, 0x99, 0xCC)
FAINT   = RGBColor(0x6F, 0x72, 0xAA)
CARD    = RGBColor(0x14, 0x14, 0x3A)
CARD2   = RGBColor(0x1B, 0x1B, 0x49)
BORDER  = RGBColor(0x2D, 0x2D, 0x68)
FRAMEBG = RGBColor(0x10, 0x10, 0x2E)
PILL    = RGBColor(0x20, 0x20, 0x4E)
INDIGO  = RGBColor(0x63, 0x66, 0xF1)
VIOLET  = RGBColor(0x8B, 0x5C, 0xF6)
CYAN    = RGBColor(0x06, 0xB6, 0xD4)
CYAN_L  = RGBColor(0x22, 0xD3, 0xEE)
IND_L   = RGBColor(0xA5, 0xB4, 0xFC)
GREEN   = RGBColor(0x34, 0xD3, 0x99)
RED     = RGBColor(0xF8, 0x71, 0x71)
AMBER   = RGBColor(0xFB, 0xBF, 0x24)
DOTR, DOTY, DOTG = RGBColor(0xF8, 0x71, 0x71), RGBColor(0xFB, 0xBF, 0x24), RGBColor(0x34, 0xD3, 0x99)

BRAND_STOPS = [(0, "6366F1"), (55, "8B5CF6"), (100, "06B6D4")]

HEAD = "Segoe UI"
BODY_F = "Segoe UI"
MONO = "Consolas"

SLIDE_W, SLIDE_H = 13.333, 7.5
MX = 0.57          # left/right margin (in)
CW = SLIDE_W - 2 * MX   # content width


# ------------------------------------------------------------------ fills ---
def _gradfill_el(stops, angle_deg):
    grad = etree.Element(qn("a:gradFill"))
    grad.set("rotWithShape", "1")
    gsLst = etree.SubElement(grad, qn("a:gsLst"))
    for pos, hexcol in stops:
        gs = etree.SubElement(gsLst, qn("a:gs"))
        gs.set("pos", str(int(pos * 1000)))
        clr = etree.SubElement(gs, qn("a:srgbClr"))
        clr.set("val", hexcol)
    lin = etree.SubElement(grad, qn("a:lin"))
    lin.set("ang", str(int(angle_deg * 60000)))
    lin.set("scaled", "1")
    return grad


def set_gradient(shape, stops=BRAND_STOPS, angle_deg=45):
    """Apply a linear gradient fill to an autoshape."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = INDIGO
    spPr = shape._element.spPr
    sf = spPr.find(qn("a:solidFill"))
    grad = _gradfill_el(stops, angle_deg)
    sf.addprevious(grad)
    spPr.remove(sf)


def run_gradient(run, stops=BRAND_STOPS, angle_deg=0):
    """Apply a gradient fill to a text run."""
    rPr = run._r.get_or_add_rPr()
    grad = _gradfill_el(stops, angle_deg)
    sf = rPr.find(qn("a:solidFill"))
    if sf is not None:
        sf.addprevious(grad)
        rPr.remove(sf)
    else:
        ln = rPr.find(qn("a:ln"))
        if ln is not None:
            ln.addnext(grad)
        else:
            rPr.insert(0, grad)


def set_shadow(shape, blur=0.09, dist=0.05, alpha=45):
    """Soft outer shadow via a:effectLst."""
    spPr = shape._element.spPr
    old = spPr.find(qn("a:effectLst"))
    if old is not None:
        spPr.remove(old)
    el = etree.Element(qn("a:effectLst"))
    sh = etree.SubElement(el, qn("a:outerShdw"))
    sh.set("blurRad", str(int(blur * 914400)))
    sh.set("dist", str(int(dist * 914400)))
    sh.set("dir", "5400000")
    sh.set("rotWithShape", "0")
    clr = etree.SubElement(sh, qn("a:srgbClr"))
    clr.set("val", "000000")
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha * 1000)))
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        ln.addnext(el)
    else:
        spPr.append(el)


def rounded_pic(pic, adj=3500):
    """Give a picture rounded corners (prstGeom roundRect)."""
    spPr = pic._element.spPr
    geom = spPr.find(qn("a:prstGeom"))
    if geom is None:
        geom = etree.SubElement(spPr, qn("a:prstGeom"))
    geom.set("prst", "roundRect")
    for child in list(geom):
        geom.remove(child)
    av = etree.SubElement(geom, qn("a:avLst"))
    gd = etree.SubElement(av, qn("a:gd"))
    gd.set("name", "adj")
    gd.set("fmla", f"val {adj}")


# -------------------------------------------------------------- transitions --
def set_transition(slide, kind="fade", direction=None):
    """Add a slide transition. advClick only - presentation stays manual."""
    sld = slide._element
    old = sld.find(qn("p:transition"))
    if old is not None:
        sld.remove(old)
    trans = etree.Element(qn("p:transition"))
    trans.set("spd", "med")
    trans.set("advClick", "1")
    child = etree.SubElement(trans, qn("p:" + kind))
    if kind == "fade":
        child.set("thruBlk", "1")
    elif kind == "split":
        child.set("orient", "horz")
        child.set("dir", direction or "out")
    elif direction:
        child.set("dir", direction)
    cSld = sld.find(qn("p:cSld"))
    siblings = list(sld)
    idx = siblings.index(cSld)
    nxt = siblings[idx + 1] if idx + 1 < len(siblings) else None
    if nxt is not None and nxt.tag == qn("p:clrMapOvr"):
        nxt.addnext(trans)
    else:
        cSld.addnext(trans)


# --------------------------------------------------------------- animations --
class Animator:
    """Collects shapes and builds one auto-start, staggered fade-in sequence.

    Animations play when the slide appears; advancing stays on mouse click /
    key press (no auto-advance), so the presenter is always in control.
    """

    def __init__(self, slide):
        self.slide = slide
        self.items = []  # (spid, delay_ms, dur_ms)

    def fade(self, *shapes, delay=0, dur=550):
        for sh in shapes:
            if sh is not None:
                self.items.append((sh.shape_id, delay, dur))
        return self

    def build(self):
        if not self.items:
            return
        counter = [100]

        def nid():
            counter[0] += 1
            return str(counter[0])

        def P(tag):
            return qn("p:" + tag)

        def cond(parent_ctn, delay):
            stc = etree.SubElement(parent_ctn, P("stCondLst"))
            c = etree.SubElement(stc, P("cond"))
            c.set("delay", str(delay))

        timing = etree.Element(P("timing"))
        tnLst = etree.SubElement(timing, P("tnLst"))
        root_par = etree.SubElement(tnLst, P("par"))
        root_ctn = etree.SubElement(root_par, P("cTn"))
        root_ctn.set("id", "1")
        root_ctn.set("dur", "indefinite")
        root_ctn.set("restart", "never")
        root_ctn.set("nodeType", "tmRoot")
        root_kids = etree.SubElement(root_ctn, P("childTnLst"))

        seq = etree.SubElement(root_kids, P("seq"))
        seq.set("concurrent", "1")
        seq.set("nextAc", "seek")
        seq_ctn = etree.SubElement(seq, P("cTn"))
        seq_ctn.set("id", "2")
        seq_ctn.set("dur", "indefinite")
        seq_ctn.set("nodeType", "mainSeq")
        seq_kids = etree.SubElement(seq_ctn, P("childTnLst"))

        step_par = etree.SubElement(seq_kids, P("par"))
        step_ctn = etree.SubElement(step_par, P("cTn"))
        step_ctn.set("id", nid())
        step_ctn.set("fill", "hold")
        cond(step_ctn, 0)
        step_kids = etree.SubElement(step_ctn, P("childTnLst"))

        for spid, delay, dur in self.items:
            holder = etree.SubElement(step_kids, P("par"))
            h_ctn = etree.SubElement(holder, P("cTn"))
            h_ctn.set("id", nid())
            h_ctn.set("fill", "hold")
            cond(h_ctn, 0)
            h_kids = etree.SubElement(h_ctn, P("childTnLst"))

            eff_par = etree.SubElement(h_kids, P("par"))
            e_ctn = etree.SubElement(eff_par, P("cTn"))
            e_ctn.set("id", nid())
            e_ctn.set("presetID", "10")
            e_ctn.set("presetClass", "entr")
            e_ctn.set("presetSubtype", "0")
            e_ctn.set("fill", "hold")
            e_ctn.set("grpId", "0")
            e_ctn.set("nodeType", "withEffect")
            cond(e_ctn, delay)
            e_kids = etree.SubElement(e_ctn, P("childTnLst"))

            # make visible, then fade in
            set_el = etree.SubElement(e_kids, P("set"))
            cb = etree.SubElement(set_el, P("cBhvr"))
            s_ctn = etree.SubElement(cb, P("cTn"))
            s_ctn.set("id", nid())
            s_ctn.set("dur", "1")
            s_ctn.set("fill", "hold")
            cond(s_ctn, 0)
            anl = etree.SubElement(s_ctn, P("attrNameLst"))
            an = etree.SubElement(anl, P("attrName"))
            an.text = "style.visibility"
            tgt = etree.SubElement(cb, P("tgtEl"))
            spt = etree.SubElement(tgt, P("spTgt"))
            spt.set("spid", str(spid))
            to = etree.SubElement(set_el, P("to"))
            sv = etree.SubElement(to, P("strVal"))
            sv.set("val", "visible")

            ae = etree.SubElement(e_kids, P("animEffect"))
            ae.set("transition", "in")
            ae.set("filter", "fade")
            cb2 = etree.SubElement(ae, P("cBhvr"))
            a_ctn = etree.SubElement(cb2, P("cTn"))
            a_ctn.set("id", nid())
            a_ctn.set("dur", str(dur))
            tgt2 = etree.SubElement(cb2, P("tgtEl"))
            spt2 = etree.SubElement(tgt2, P("spTgt"))
            spt2.set("spid", str(spid))

        for name in ("prevCondLst", "nextCondLst"):
            cl = etree.SubElement(seq, P(name))
            c = etree.SubElement(cl, P("cond"))
            c.set("delay", "indefinite")
            c.set("evt", "onClick")
            t = etree.SubElement(c, P("tgtEl"))
            etree.SubElement(t, P("sldTgt"))

        self.slide._element.append(timing)


# ------------------------------------------------------------------- text ---
def tb(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return box, tf


def para(tf, first=False, align=PP_ALIGN.LEFT, space_before=0, space_after=0,
         line=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if line:
        p.line_spacing = line
    return p


def run(p, text, size=13, color=BODY, bold=False, font=BODY_F, italic=False,
        spc=None, gradient=False):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = color
    if spc is not None:
        r._r.get_or_add_rPr().set("spc", str(int(spc * 100)))
    if gradient:
        run_gradient(r)
    return r


def _hanging(p, indent_in=0.19):
    pPr = p._p.get_or_add_pPr()
    emu = str(int(indent_in * 914400))
    pPr.set("marL", emu)
    pPr.set("indent", "-" + emu)


def bullets(slide, x, y, w, h, items, size=12.5, gap=7, color=BODY,
            marker="●", marker_color=CYAN_L, line=1.12, anchor=MSO_ANCHOR.TOP,
            marker_scale=0.62):
    """items: strings, or (lead, rest) tuples -> lead rendered bold white."""
    box, tf = tb(slide, x, y, w, h, anchor=anchor)
    for i, item in enumerate(items):
        p = para(tf, first=(i == 0), space_after=gap, line=line)
        _hanging(p)
        run(p, marker + "  ", size=size * marker_scale, color=marker_color,
            bold=True)
        if isinstance(item, tuple):
            run(p, item[0], size=size, color=WHITE, bold=True)
            run(p, item[1], size=size, color=color)
        else:
            run(p, item, size=size, color=color)
    return box


# ----------------------------------------------------------------- shapes ---
def card(slide, x, y, w, h, fill=CARD, line=BORDER, radius=0.06, line_w=1.0,
         shadow=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    if shadow:
        set_shadow(sh)
    sh.text_frame.paragraphs[0].text = ""
    return sh


def oval(slide, x, y, d, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                                Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def grad_bar(slide, x, y, w, h=0.07, angle=0):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        sh.adjustments[0] = 0.5
    except Exception:
        pass
    sh.line.fill.background()
    sh.shadow.inherit = False
    set_gradient(sh, angle_deg=angle)
    return sh


def chip(slide, x, y, w, h, text, size=10.5, fill=PILL, color=BODY,
         bold=False, align=PP_ALIGN.CENTER):
    sh = card(slide, x, y, w, h, fill=fill, line=BORDER, radius=0.5, line_w=0.75)
    tf = sh.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run(p, text, size=size, color=color, bold=bold)
    return sh


# ------------------------------------------------------------ image frames --
def browser(slide, img, x, y, w, ratio=1.60, url="quizmaster.app", bar_h=0.34):
    """Screenshot inside a browser-chrome frame. Returns (shapes, total_h).
    The image was pre-cropped to `ratio`, so it fits with zero distortion."""
    pad = 0.055
    iw = w - 2 * pad
    ih = iw / ratio
    total_h = bar_h + ih + pad
    shapes = []
    frame = card(slide, x, y, w, total_h, fill=FRAMEBG, line=BORDER,
                 radius=0.035, line_w=1.1)
    set_shadow(frame, blur=0.14, dist=0.06, alpha=50)
    shapes.append(frame)
    for i, c in enumerate((DOTR, DOTY, DOTG)):
        shapes.append(oval(slide, x + 0.17 + i * 0.21, y + bar_h / 2 - 0.05,
                           0.10, c))
    pw = min(w * 0.46, 3.1)
    pill = card(slide, x + (w - pw) / 2, y + bar_h / 2 - 0.105, pw, 0.21,
                fill=PILL, line=None, radius=0.5)
    shapes.append(pill)
    tbox, tf = tb(slide, x + (w - pw) / 2, y + bar_h / 2 - 0.105, pw, 0.21,
                  anchor=MSO_ANCHOR.MIDDLE)
    p = para(tf, first=True, align=PP_ALIGN.CENTER)
    run(p, url, size=8.5, color=MUTED, font=BODY_F)
    shapes.append(tbox)
    pic = slide.shapes.add_picture(img, Inches(x + pad), Inches(y + bar_h),
                                   Inches(iw), Inches(ih))
    rounded_pic(pic, 2600)
    shapes.append(pic)
    return shapes, total_h


def phone(slide, img, x, y, w, ratio=0.52):
    """Screenshot inside a phone frame. Returns (shapes, total_h)."""
    pad = 0.055
    iw = w - 2 * pad
    ih = iw / ratio
    total_h = ih + 2 * pad + 0.05
    shapes = []
    frame = card(slide, x, y, w, total_h, fill=FRAMEBG,
                 line=RGBColor(0x3D, 0x3D, 0x7D), radius=0.085, line_w=1.4)
    set_shadow(frame, blur=0.14, dist=0.06, alpha=55)
    shapes.append(frame)
    pic = slide.shapes.add_picture(img, Inches(x + pad), Inches(y + pad + 0.03),
                                   Inches(iw), Inches(ih))
    rounded_pic(pic, 5200)
    shapes.append(pic)
    notch = card(slide, x + w / 2 - 0.30, y + 0.10, 0.60, 0.085,
                 fill=RGBColor(0x05, 0x05, 0x16), line=None, radius=0.5)
    shapes.append(notch)
    return shapes, total_h


# --------------------------------------------------------------- scaffold ---
def add_bg(slide, img):
    return slide.shapes.add_picture(img, 0, 0, Inches(SLIDE_W),
                                    Inches(SLIDE_H))


def footer(slide, page, total=11):
    _, tf = tb(slide, MX, 7.12, 7.5, 0.3)
    p = para(tf, first=True)
    run(p, "QuizMaster", size=9, color=IND_L, bold=True)
    run(p, "  ·  Learn Smarter · Assess Better · Improve Continuously",
        size=9, color=FAINT)
    _, tf2 = tb(slide, SLIDE_W - MX - 1.6, 7.12, 1.6, 0.3)
    p2 = para(tf2, first=True, align=PP_ALIGN.RIGHT)
    run(p2, f"{page:02d}", size=10, color=WHITE, bold=True)
    run(p2, f" / {total:02d}", size=10, color=FAINT)


def header(slide, kicker, title_runs, subtitle=None, title_size=27):
    """Kicker + big title (+ optional subtitle). Returns shapes for animation.
    title_runs: list of (text, use_gradient)."""
    shapes = []
    bar = grad_bar(slide, MX, 0.47, 0.52, 0.075)
    shapes.append(bar)
    kb, tf = tb(slide, MX + 0.66, 0.345, 9.5, 0.32)
    p = para(tf, first=True)
    run(p, kicker.upper(), size=11.5, color=CYAN_L, bold=True, spc=2.6)
    shapes.append(kb)
    tbox, tf2 = tb(slide, MX, 0.66, CW, 0.75)
    p2 = para(tf2, first=True, line=1.02)
    for text, grad in title_runs:
        run(p2, text, size=title_size, color=WHITE, bold=True, font=HEAD,
            gradient=grad)
    shapes.append(tbox)
    if subtitle:
        sb, tf3 = tb(slide, MX, 1.30, CW, 0.4)
        p3 = para(tf3, first=True)
        run(p3, subtitle, size=13, color=MUTED)
        shapes.append(sb)
    return shapes


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text
